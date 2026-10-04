"""Decode bounded Office display text using the shipped independent Python libraries."""
import csv
import datetime
import html
import json
import os
import sys
import zipfile

MAX_ROWS = 1000
MAX_COLUMNS = 100
MAX_SHEETS = 20
MAX_TEXT = 500_000
MAX_CELL = 4000


def archive_check(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError("Office archive exceeds the 64 MiB expanded preview limit")
        if any(entry.file_size > 16 * 1024 * 1024 for entry in entries):
            raise ValueError("Office archive member exceeds the 16 MiB preview limit")


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat()
    return str(value)


def sheets_preview(path, extension):
    truncated = False
    sheets = []
    characters = 0

    def rows_preview(rows):
        nonlocal truncated, characters
        output = []
        for index, row in enumerate(rows):
            if index >= MAX_ROWS or characters >= MAX_TEXT:
                truncated = True
                break
            if len(row) > MAX_COLUMNS:
                truncated = True
            cells = []
            for value in row[:MAX_COLUMNS]:
                cell = text(value)
                allowed = min(MAX_CELL, max(0, MAX_TEXT - characters))
                if len(cell) > allowed:
                    truncated = True
                cells.append(cell[:allowed])
                characters += min(len(cell), allowed)
            output.append(cells)
        return output

    if extension == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > MAX_SHEETS:
                truncated = True
            for sheet in workbook.worksheets[:MAX_SHEETS]:
                row_count = sheet.max_row or MAX_ROWS + 1
                column_count = sheet.max_column or MAX_COLUMNS + 1
                if row_count > MAX_ROWS or column_count > MAX_COLUMNS:
                    truncated = True
                rows = sheet.iter_rows(max_row=min(row_count, MAX_ROWS), max_col=min(column_count, MAX_COLUMNS), values_only=True)
                sheets.append({"name": sheet.title, "rows": rows_preview(rows)})
        finally:
            workbook.close()
    else:
        csv.field_size_limit(1024 * 1024)
        with open(path, encoding="utf-8-sig", newline="") as source:
            sheets.append({"name": os.path.basename(path), "rows": rows_preview(csv.reader(source, delimiter="\t" if extension == ".tsv" else ","))})
    return {"kind": "spreadsheet", "sheets": sheets, "truncated": truncated}


def document_preview(path, extension):
    pieces = []
    characters = 0
    truncated = False

    def escape(value):
        nonlocal characters, truncated
        value = text(value)
        allowed = max(0, MAX_TEXT - characters)
        characters += min(allowed, len(value))
        truncated = truncated or len(value) > allowed
        return html.escape(value[:allowed]).replace("\n", "<br>")

    if extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        for block in document.iter_inner_content():
            if characters >= MAX_TEXT:
                truncated = True
                break
            if isinstance(block, Paragraph):
                style = block.style.name if block.style else ""
                tag = "h" + style[-1] if style.startswith("Heading ") and style[-1:] in "123456" else "p"
                pieces.append("<" + tag + ">" + escape(block.text) + "</" + tag + ">")
            elif isinstance(block, Table):
                rows = []
                if len(block.rows) > MAX_ROWS:
                    truncated = True
                for row in block.rows[:MAX_ROWS]:
                    if len(row.cells) > MAX_COLUMNS:
                        truncated = True
                    rows.append("<tr>" + "".join("<td>" + escape(cell.text) + "</td>" for cell in row.cells[:MAX_COLUMNS]) + "</tr>")
                pieces.append("<table>" + "".join(rows) + "</table>")
    else:
        from pptx import Presentation
        slides = Presentation(path).slides
        if len(slides) > 100:
            truncated = True
        for number, slide in enumerate(slides):
            if number >= 100 or characters >= MAX_TEXT:
                truncated = True
                break
            content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    content.append("<p>" + escape(shape.text) + "</p>")
                if shape.has_table:
                    rows = shape.table.rows
                    truncated = truncated or len(rows) > MAX_ROWS
                    table_rows = []
                    for row in list(rows)[:MAX_ROWS]:
                        cells = list(row.cells)
                        truncated = truncated or len(cells) > MAX_COLUMNS
                        table_rows.append("<tr>" + "".join("<td>" + escape(cell.text) + "</td>" for cell in cells[:MAX_COLUMNS]) + "</tr>")
                    content.append("<table>" + "".join(table_rows) + "</table>")
            pieces.append('<section aria-label="' + str(number + 1) + '">' + "".join(content) + "</section>")
    style = "body{font:16px system-ui;margin:24px;line-height:1.6;color:#20242b;background:#fff;overflow-wrap:anywhere}table{border-collapse:collapse;max-width:100%}td{border:1px solid #ccd2da;padding:6px 10px;white-space:pre-wrap}p{white-space:pre-wrap}section{border-bottom:1px solid #ccd2da;padding:16px 0}h1,h2,h3{line-height:1.25}"
    output = '<!doctype html><html><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'"><style>' + style + '</style></head><body>' + "".join(pieces) + '</body></html>'
    return {"kind": "document", "html": output, "truncated": truncated}


def main():
    path = sys.argv[1]
    if os.path.getsize(path) > 32 * 1024 * 1024:
        raise ValueError("Office file exceeds the 32 MiB preview limit")
    extension = os.path.splitext(path)[1].lower()
    if extension in (".xlsx", ".docx", ".pptx"):
        archive_check(path)
    if extension in (".xlsx", ".csv", ".tsv"):
        result = sheets_preview(path, extension)
    elif extension in (".docx", ".pptx"):
        result = document_preview(path, extension)
    else:
        raise ValueError("Unsupported Office preview extension")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
